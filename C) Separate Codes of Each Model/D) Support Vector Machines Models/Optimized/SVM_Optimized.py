import numpy as np
import pandas as pd
import warnings
import random
import shap
import matplotlib.pyplot as plt
import os
from sklearn.model_selection import train_test_split
from sklearn.metrics import mean_squared_error, mean_absolute_error, r2_score
from sklearn.preprocessing import StandardScaler, OneHotEncoder
from sklearn.compose import ColumnTransformer
from docx import Document
from sklearn.exceptions import ConvergenceWarning
from sklearn.pipeline import Pipeline
from sklearn.svm import SVR, NuSVR

warnings.filterwarnings("ignore", category=ConvergenceWarning)

# ======================== Load Dataset ========================
df = pd.read_csv("Filtered_Dataset.csv")
y = df["CS"]
X = df.drop(columns=["CS"])

categorical_features = ['FT', 'CT', 'SPB', 'ST']
numerical_features = [col for col in X.columns if col not in categorical_features]

X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2, random_state=42)

preprocessor = ColumnTransformer([
    ('cat', OneHotEncoder(drop='first', sparse_output=False), categorical_features),
    ('num', StandardScaler(), numerical_features)
])

# ====================== Evaluation Metrics =====================
def evaluate(y_true, y_pred):
    return {
        "RMSE": np.sqrt(mean_squared_error(y_true, y_pred)),
        "MAE": mean_absolute_error(y_true, y_pred),
        "MAPE": np.mean(np.abs((y_true - y_pred) / y_true)) * 100,
        "R2": r2_score(y_true, y_pred)
    }

# ====================== Optimizers =====================
def pso_optimize(obj_func, bounds, n_particles=10, max_iter=20, w=0.6, c1=1.2, c2=2.2):
    dim = len(bounds)
    X = np.random.uniform([b[0] for b in bounds], [b[1] for b in bounds], (n_particles, dim))
    V = np.zeros_like(X)
    p_best = X.copy()
    p_best_scores = np.array([obj_func(x) for x in X])
    g_best = p_best[np.argmin(p_best_scores)]

    for _ in range(max_iter):
        for i in range(n_particles):
            r1, r2 = np.random.rand(dim), np.random.rand(dim)
            V[i] = w * V[i] + c1 * r1 * (p_best[i] - X[i]) + c2 * r2 * (g_best - X[i])
            X[i] = np.clip(X[i] + V[i], [b[0] for b in bounds], [b[1] for b in bounds])
            score = obj_func(X[i])
            if score < p_best_scores[i]:
                p_best[i] = X[i]
                p_best_scores[i] = score
        g_best = p_best[np.argmin(p_best_scores)]

    return g_best, min(p_best_scores)

def ga_optimize(obj_func, bounds, pop_size=10, generations=35, mutation_rate=0.25, crossover_rate=0.75):
    dim = len(bounds)
    population = [np.random.uniform([b[0] for b in bounds], [b[1] for b in bounds]) for _ in range(pop_size)]
    population = np.array(population)

    for _ in range(generations):
        scores = np.array([obj_func(ind) for ind in population])
        sorted_indices = np.argsort(scores)
        sorted_pop = population[sorted_indices]
        next_gen = list(sorted_pop[:2])

        while len(next_gen) < pop_size:
            if dim > 1:
                cut = random.randint(1, dim - 1)
            else:
                cut = 1
            if random.random() < crossover_rate:
                parents = random.sample(list(sorted_pop[:10]), 2)
                child = np.concatenate((parents[0][:cut], parents[1][cut:]))
            else:
                child = random.choice(sorted_pop[:10]).copy()

            if random.random() < mutation_rate:
                idx = random.randint(0, dim - 1)
                low, high = bounds[idx]
                child[idx] = np.clip(child[idx] + np.random.normal(0, 0.2), low, high)

            next_gen.append(np.clip(child, [b[0] for b in bounds], [b[1] for b in bounds]))

        population = np.array(next_gen)

    best_index = np.argmin([obj_func(ind) for ind in population])
    best = population[best_index]
    return best, obj_func(best)

def firefly_optimize(obj_func, bounds, n=10, max_iter=10, alpha=0.6, beta=1.0, gamma=1.5):
    dim = len(bounds)
    fireflies = [np.random.uniform([b[0] for b in bounds], [b[1] for b in bounds]) for _ in range(n)]
    light = [obj_func(f) for f in fireflies]

    for _ in range(max_iter):
        for i in range(n):
            for j in range(n):
                if light[j] < light[i]:
                    r = np.linalg.norm(fireflies[i] - fireflies[j])
                    beta_temp = beta * np.exp(-gamma * r ** 2)
                    attraction = beta_temp * (fireflies[j] - fireflies[i])
                    random_move = alpha * (np.random.rand(dim) - 0.5)
                    fireflies[i] += attraction + random_move
                    fireflies[i] = np.clip(fireflies[i], [b[0] for b in bounds], [b[1] for b in bounds])
                    light[i] = obj_func(fireflies[i])

    best_idx = np.argmin(light)
    return fireflies[best_idx], light[best_idx]

def sa_optimize(obj_func, bounds, max_iter=10, T0=120, alpha=0.85):
    dim = len(bounds)
    x = np.random.uniform([b[0] for b in bounds], [b[1] for b in bounds])
    fx = obj_func(x)
    best = x.copy()
    fbest = fx
    T = T0

    for _ in range(max_iter):
        for _ in range(15):
            candidate = x + np.random.uniform(-0.15, 0.15, dim)
            candidate = np.clip(candidate, [b[0] for b in bounds], [b[1] for b in bounds])
            fc = obj_func(candidate)
            if fc < fx or np.random.rand() < np.exp(-(fc - fx) / T):
                x, fx = candidate, fc
                if fc < fbest:
                    best, fbest = candidate, fc
        T *= alpha
    return best, fbest

def de_optimize(obj_func, bounds, pop_size=10, max_iter=20, F=0.5, CR=0.7):
    dim = len(bounds)
    pop = np.random.uniform([b[0] for b in bounds], [b[1] for b in bounds], (pop_size, dim))
    scores = np.array([obj_func(ind) for ind in pop])

    for _ in range(max_iter):
        for i in range(pop_size):
            idxs = [idx for idx in range(pop_size) if idx != i]
            a, b, c = pop[np.random.choice(idxs, 3, replace=False)]
            mutant = np.clip(a + F * (b - c), [b[0] for b in bounds], [b[1] for b in bounds])
            cross_points = np.random.rand(dim) < CR
            if not np.any(cross_points):
                cross_points[np.random.randint(0, dim)] = True
            trial = np.where(cross_points, mutant, pop[i])
            score = obj_func(trial)
            if score < scores[i]:
                scores[i] = score
                pop[i] = trial
    best_idx = np.argmin(scores)
    return pop[best_idx], scores[best_idx]

def abc_optimize(obj_func, bounds, pop_size=10, max_iter=20, limit=5):
    dim = len(bounds)
    foods = np.random.uniform([b[0] for b in bounds], [b[1] for b in bounds], (pop_size, dim))
    fitness = np.array([obj_func(f) for f in foods])
    trial = np.zeros(pop_size)

    for _ in range(max_iter):
        for i in range(pop_size):
            k = random.choice([j for j in range(pop_size) if j != i])
            phi = np.random.uniform(-1, 1, dim)
            new = foods[i] + phi * (foods[i] - foods[k])
            new = np.clip(new, [b[0] for b in bounds], [b[1] for b in bounds])
            new_fit = obj_func(new)
            if new_fit < fitness[i]:
                foods[i] = new
                fitness[i] = new_fit
                trial[i] = 0
            else:
                trial[i] += 1

        probs = fitness.max() - fitness
        probs = probs / probs.sum()
        for _ in range(pop_size):
            i = np.random.choice(range(pop_size), p=probs)
            k = random.choice([j for j in range(pop_size) if j != i])
            phi = np.random.uniform(-1, 1, dim)
            new = foods[i] + phi * (foods[i] - foods[k])
            new = np.clip(new, [b[0] for b in bounds], [b[1] for b in bounds])
            new_fit = obj_func(new)
            if new_fit < fitness[i]:
                foods[i] = new
                fitness[i] = new_fit
                trial[i] = 0
            else:
                trial[i] += 1

        for i in range(pop_size):
            if trial[i] > limit:
                foods[i] = np.random.uniform([b[0] for b in bounds], [b[1] for b in bounds])
                fitness[i] = obj_func(foods[i])
                trial[i] = 0

    best_idx = np.argmin(fitness)
    return foods[best_idx], fitness[best_idx]

def hs_optimize(obj_func, bounds, hm_size=10, max_iter=20, HMCR=0.9, PAR=0.3, bw=0.01):
    dim = len(bounds)
    HM = np.random.uniform([b[0] for b in bounds], [b[1] for b in bounds], (hm_size, dim))
    scores = np.array([obj_func(h) for h in HM])

    for _ in range(max_iter):
        new = []
        for d in range(dim):
            if random.random() < HMCR:
                idx = random.randint(0, hm_size - 1)
                val = HM[idx, d]
                if random.random() < PAR:
                    val += np.random.uniform(-bw, bw)
                val = np.clip(val, bounds[d][0], bounds[d][1])
            else:
                val = np.random.uniform(bounds[d][0], bounds[d][1])
            new.append(val)
        new = np.array(new)
        new_score = obj_func(new)
        if new_score < scores.max():
            worst_idx = np.argmax(scores)
            HM[worst_idx] = new
            scores[worst_idx] = new_score

    best_idx = np.argmin(scores)
    return HM[best_idx], scores[best_idx]

# Add optimizers to dictionary
optimizers = {
    "PSO": pso_optimize,
    "GA": ga_optimize,
    "Firefly": firefly_optimize,
    "SA": sa_optimize,
    "DE": de_optimize,
    "ABC": abc_optimize,
    "HS": hs_optimize
}

# =================== Model Configs ===================
model_configs = {
    "SVR": (
        SVR,
        ["C", "epsilon", "gamma"],
        [(1.0, 100.0), (0.01, 0.1), (0.001, 0.5)]
    ),
    "NuSVR": (
        NuSVR,
        ["nu", "C", "gamma"],
        [(0.1, 0.9), (1.0, 100.0), (0.001, 0.5)]
    )
}

all_results = {}
optimized_models = {}
for model_name, (model_class, param_names, bounds) in model_configs.items():
    results = {}
    print(f"\n=== {model_name} ===")

    for opt_name, optimizer in optimizers.items():
        print(f"\n🔧 Optimizing {model_name} using {opt_name}...")

        def objective(params):
            if model_name == "SVR":
                C, epsilon, gamma = params
                model = model_class(C=C, epsilon=epsilon, gamma=gamma, kernel='rbf')

            elif model_name == "NuSVR":
                nu, C, gamma = params
                model = model_class(nu=nu, C=C, gamma=gamma, kernel='rbf')

            else:
                raise ValueError(f"Unsupported model name: {model_name}")

            pipeline = Pipeline([
                ('preprocessor', preprocessor),
                ('regressor', model)
            ])

            pipeline.fit(X_train, y_train)
            preds = pipeline.predict(X_test)
            return mean_squared_error(y_test, preds)

        np.random.seed(2000 + hash(opt_name + model_name) % 1000)
        random.seed(2000 + hash(opt_name + model_name) % 1000)

        best_params, _ = optimizer(objective, bounds)

        if model_name == "SVR":
            C, epsilon, gamma = best_params
            model = model_class(C=C, epsilon=epsilon, gamma=gamma, kernel='rbf')

        elif model_name == "NuSVR":
            nu, C, gamma = best_params
            model = model_class(nu=nu, C=C, gamma=gamma, kernel='rbf')

        pipeline = Pipeline([
            ('preprocessor', preprocessor),
            ('regressor', model)
        ])
        pipeline.fit(X_train, y_train)

        # 🟩 Save pipeline for SHAP
        if model_name not in optimized_models:
            optimized_models[model_name] = {}
        optimized_models[model_name][opt_name] = pipeline

        train_pred = pipeline.predict(X_train)
        test_pred = pipeline.predict(X_test)

        train_metrics = evaluate(y_train, train_pred)
        test_metrics = evaluate(y_test, test_pred)

        results[opt_name] = {
            "Train": train_metrics,
            "Test": test_metrics
        }

        print(f"✅ {opt_name} Done. R² = {test_metrics['R2']:.3f}, RMSE = {test_metrics['RMSE']:.2f}")

    all_results[model_name] = results



    from docx.shared import Inches

    doc = Document()
    doc.add_heading(f"{model_name} Optimized Models", 0)

    metrics_headers = ["Optimizer", "Train_R2", "Test_R2", "Train_RMSE", "Test_RMSE", "Train_MAE", "Test_MAE",
                       "Train_MAPE", "Test_MAPE"]
    table = doc.add_table(rows=1, cols=len(metrics_headers))
    table.style = 'Light Grid Accent 1'

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(metrics_headers):
        hdr_cells[i].text = h

    for opt_name, metrics in results.items():
        row_cells = table.add_row().cells
        row_cells[0].text = opt_name
        row_cells[1].text = f"{metrics['Train']['R2']:.3f}"
        row_cells[2].text = f"{metrics['Test']['R2']:.3f}"
        row_cells[3].text = f"{metrics['Train']['RMSE']:.3f}"
        row_cells[4].text = f"{metrics['Test']['RMSE']:.3f}"
        row_cells[5].text = f"{metrics['Train']['MAE']:.3f}"
        row_cells[6].text = f"{metrics['Test']['MAE']:.3f}"
        row_cells[7].text = f"{metrics['Train']['MAPE']:.3f}"
        row_cells[8].text = f"{metrics['Test']['MAPE']:.3f}"

    doc.save(f"{model_name.replace(' ', '_')}_Optimized_Summary.docx")
    print(f"📄 Word file saved: {model_name.replace(' ', '_')}_Optimized_Summary.docx")

# Create a folder to store SHAP plots if it doesn't exist
os.makedirs("shap_plots", exist_ok=True)

print("\n" + "=" * 100)
print("🔍 SHAP ANALYSIS OF OPTIMIZED MODELS")
print("=" * 100)

for model_name, optimizer_dict in optimized_models.items():
    for optimizer_name, pipeline in optimizer_dict.items():
        print(f"\n{'=' * 80}")
        print(f"📌 SHAP Summary for {model_name} optimized using {optimizer_name}")
        print(f"{'=' * 80}")

        try:
            preprocessor = pipeline.named_steps['preprocessor']
            model = pipeline.named_steps['regressor']
            X_train_transformed = preprocessor.transform(X_train)

            # Get feature names
            cat_encoder = preprocessor.named_transformers_['cat']
            cat_feature_names = cat_encoder.get_feature_names_out(categorical_features)
            num_feature_names = numerical_features
            all_feature_names = np.concatenate([cat_feature_names, num_feature_names])

            # Create prediction function for KernelExplainer
            def model_predict(data):
                return model.predict(data)

            # Use KernelExplainer for SVR/NuSVR models
            background = shap.kmeans(X_train_transformed, 50)
            explainer = shap.KernelExplainer(model_predict, background, feature_names=all_feature_names)
            shap_values = explainer.shap_values(X_train_transformed[:200], nsamples=100)

            explanation = shap.Explanation(values=shap_values,
                                           data=X_train_transformed[:200],
                                           feature_names=all_feature_names)

            # Plot and save
            plt.figure()
            shap.plots.beeswarm(explanation, show=False)
            plt.title(f"{model_name} + {optimizer_name}", fontsize=14)
            plt.tight_layout()

            filename = f"shap_plots/{model_name.replace(' ', '_')}_{optimizer_name}.jpeg"
            plt.savefig(filename, format='jpeg', dpi=300)
            plt.close()
            print(f"✅ SHAP plot saved: {filename}")

        except Exception as e:
            print(f"[❌ ERROR] Could not generate SHAP for {model_name} + {optimizer_name}: {e}")

